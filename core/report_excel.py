#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
qkl_report.py - Xuat file Excel ket qua ra soat ho so khoi luong theo mau TEXO.
Dung: python3 qkl_report.py <ket_qua.json> <output.xlsx>
Cau truc JSON: xem SKILL.md muc 7.
"""
import json, sys, datetime
from collections import Counter, OrderedDict
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.series import DataPoint
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

FONT = "Times New Roman"
NAVY = "1F3864"
LEVELS = OrderedDict([
    ("Cao",        {"fill": "F4CCCC", "font": "9C0006", "bar": "C0392B"}),
    ("Trung bình", {"fill": "FFE699", "font": "7F6000", "bar": "E6A117"}),
    ("Thấp",       {"fill": "E2EFDA", "font": "375623", "bar": "70AD47"}),
])
COMPANY_1 = "CÔNG TY CỔ PHẦN TEXO"
COMPANY_2 = "TƯ VẤN VÀ ĐẦU TƯ"
DEPT = "PHÒNG KỸ THUẬT"

thin = Side(style="thin", color="8C8C8C")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(size=11, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def al(h="left", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def setup_page(ws, paper, landscape=True, title_rows=None, fit_height=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if paper == "A3" else ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if fit_height else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.6
    ws.page_margins.header = ws.page_margins.footer = 0.3
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddFooter.left.text = "Công ty Cổ phần TEXO Tư vấn và Đầu tư – Phòng Kỹ thuật"
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.font = FONT
    ws.oddFooter.right.text = "Trang &P/&N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.font = FONT
    ws.sheet_view.showGridLines = False


def letterhead(ws, last_col, date_text, title, subtitle=None):
    """Tieu de cong ty (trai) + ngay (phai) + ten bieu."""
    mid = max(3, last_col // 2)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=mid)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=mid)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=mid)
    ws.cell(1, 1, COMPANY_1).font = f(12, True, NAVY)
    ws.cell(2, 1, COMPANY_2).font = f(12, True, NAVY)
    ws.cell(3, 1, DEPT).font = f(11, True)
    for r in (1, 2, 3):
        ws.cell(r, 1).alignment = al("center", wrap=False)
    ws.merge_cells(start_row=3, start_column=mid + 1, end_row=3, end_column=last_col)
    c = ws.cell(3, mid + 1, date_text)
    c.font = f(11, italic=True)
    c.alignment = al("right", wrap=False)
    # duong ke duoi tieu de cong ty
    for col in range(1, last_col + 1):
        ws.cell(4, col).border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[4].height = 6
    ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=last_col)
    t = ws.cell(6, 1, title)
    t.font = f(16, True, NAVY)
    t.alignment = al("center", wrap=False)
    ws.row_dimensions[6].height = 26
    nxt = 7
    if subtitle:
        ws.merge_cells(start_row=7, start_column=1, end_row=7, end_column=last_col)
        s = ws.cell(7, 1, subtitle)
        s.font = f(11, italic=True)
        s.alignment = al("center", wrap=False)
        nxt = 8
    return nxt + 1


def vn_date(iso):
    try:
        d = datetime.date.fromisoformat(iso)
    except Exception:
        d = datetime.date.today()
    return d, f"Hà Nội, ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"


def fit_rows(ws, first, last, widths, min_h=20):
    """Uoc luong chieu cao dong theo do dai chu (LibreOffice/Excel khong tu gian dong da merge)."""
    import math
    for r in range(first, last + 1):
        lines = 1
        for j, w in enumerate(widths, 1):
            v = ws.cell(r, j).value
            if v is None:
                continue
            txt = str(v)
            n = sum(max(1, math.ceil(len(part) / max(w * 1.05, 1))) for part in txt.split("\n"))
            lines = max(lines, n)
        ws.row_dimensions[r].height = max(min_h, 15 * lines + 4)


def no_extra_labels(lbls):
    lbls.showVal = True
    lbls.showCatName = False
    lbls.showSerName = False
    lbls.showLegendKey = False
    lbls.showPercent = False
    return lbls


def table_header(ws, row, headers, widths=None, height=32):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row, i, h)
        c.font = f(11, True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = al("center")
        c.border = BOX
        if widths:
            ws.column_dimensions[L(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = height


def build(data, out):
    paper = (data.get("kho_giay") or "A4").upper()
    d, date_text = vn_date(data.get("ngay", ""))
    diem = data.get("diem", [])
    order = {k: i for i, k in enumerate(LEVELS)}
    diem = sorted(diem, key=lambda x: (order.get(x.get("muc"), 9), -(x.get("ty_trong") or 0)))
    thieu = data.get("thieu_ho_so", [])
    has_tt = any(x.get("ty_trong") not in (None, "") for x in diem)
    has_mq = any(x.get("ma_quy_tac") for x in diem)

    wb = Workbook()

    # ---------------- Sheet 2: Danh sach diem (tao truoc de cong thuc tham chieu) -------------
    ds = wb.active
    ds.title = "Danh sách điểm"
    heads = ["STT", "Mức rủi ro", "Vị trí trong hồ sơ", "Phần công trình", "Hạng mục",
             "Khối lượng đề nghị", "Vấn đề cần làm rõ", "Câu hỏi cho người lập / người kiểm",
             "Cách kiểm chứng", "Hồ sơ còn thiếu"]
    widths = [6, 11, 16, 15, 30, 15, 42, 40, 32, 24]
    if has_tt:
        heads.insert(6, "Tỷ trọng (%)")
        widths.insert(6, 10)
    if has_mq:
        heads.append("Quy tắc áp dụng")
        widths.append(14)
    last = len(heads)
    r0 = letterhead(ds, last, date_text, "DANH SÁCH CÁC ĐIỂM CẦN RÀ SOÁT",
                    f"Dự án: {data.get('du_an','')}   |   Hồ sơ: {data.get('ho_so','')}")
    table_header(ds, r0, heads, widths)
    first = r0 + 1
    for i, x in enumerate(diem, 1):
        r = r0 + i
        vals = [i, x.get("muc", ""), x.get("vi_tri", ""), x.get("phan", ""), x.get("hang_muc", ""),
                x.get("khoi_luong", "")]
        if has_tt:
            vals.append(x.get("ty_trong"))
        vals += [x.get("van_de", ""), x.get("cau_hoi", ""), x.get("kiem_chung", ""), x.get("ho_so_thieu", "")]
        if has_mq:
            vals.append(x.get("ma_quy_tac", ""))
        for j, v in enumerate(vals, 1):
            c = ds.cell(r, j, v)
            c.font = f(11)
            c.border = BOX
            c.alignment = al("center" if j in (1, 2, 3, 6) or (has_tt and j == 7) else "left", "top")
            if i % 2 == 0:
                c.fill = fill("F2F5FA")
        lv = LEVELS.get(x.get("muc"))
        if lv:
            c = ds.cell(r, 2)
            c.fill = fill(lv["fill"])
            c.font = f(11, True, lv["font"])
        if has_tt:
            ds.cell(r, 7).number_format = "0.0"
    lastrow = r0 + max(len(diem), 1)
    fit_rows(ds, first, lastrow, widths)
    ds.freeze_panes = ds.cell(first, 3)
    ds.auto_filter.ref = f"A{r0}:{L(last)}{lastrow}"
    setup_page(ds, paper, title_rows=f"{r0}:{r0}")
    lvl_rng = f"'Danh sách điểm'!$B${first}:$B${max(lastrow, first)}"
    part_rng = f"'Danh sách điểm'!$D${first}:$D${max(lastrow, first)}"

    # ---------------- Sheet: Thieu ho so -------------
    th = wb.create_sheet("Thiếu hồ sơ")
    h2 = ["STT", "Hạng mục", "Khối lượng", "Hồ sơ còn thiếu", "Đơn vị bổ sung", "Hướng xử lý đề xuất"]
    w2 = [6, 38, 16, 40, 20, 40]
    r2 = letterhead(th, len(h2), date_text, "DANH SÁCH HẠNG MỤC THIẾU HỒ SƠ",
                    "Cần xin ý kiến các bên thống nhất cách xác nhận")
    table_header(th, r2, h2, w2)
    for i, x in enumerate(thieu, 1):
        vals = [i, x.get("hang_muc", ""), x.get("khoi_luong", ""), x.get("thieu", ""),
                x.get("don_vi_bo_sung", ""), x.get("huong_xu_ly", "")]
        for j, v in enumerate(vals, 1):
            c = th.cell(r2 + i, j, v)
            c.font = f(11)
            c.border = BOX
            c.alignment = al("center" if j in (1, 3) else "left", "top")
    fit_rows(th, r2 + 1, r2 + len(thieu), w2)
    if not thieu:
        th.merge_cells(start_row=r2 + 1, start_column=1, end_row=r2 + 1, end_column=len(h2))
        c = th.cell(r2 + 1, 1, "Chưa ghi nhận hạng mục thiếu hồ sơ.")
        c.font = f(11, italic=True)
        c.alignment = al("center")
    th.freeze_panes = th.cell(r2 + 1, 1)
    setup_page(th, paper, title_rows=f"{r2}:{r2}")

    # ---------------- Sheet: Giao viec -------------
    gv = wb.create_sheet("Giao việc")
    h3 = ["STT", "Mức", "Nội dung cần làm rõ", "Vị trí", "Người xử lý", "Hạn trả lời", "Kết quả", "Ghi chú"]
    w3 = [6, 11, 60, 16, 18, 13, 16, 26]
    r3 = letterhead(gv, len(h3), date_text, "PHIẾU GIAO VIỆC LÀM RÕ KHỐI LƯỢNG",
                    f"Dự án: {data.get('du_an','')}")
    table_header(gv, r3, h3, w3)
    dv = DataValidation(type="list", formula1='"Đã làm rõ,Cần bổ sung hồ sơ,Điều chỉnh khối lượng,Chưa xử lý"',
                        allow_blank=True)
    gv.add_data_validation(dv)
    for i, x in enumerate(diem, 1):
        r = r3 + i
        vals = [i, x.get("muc", ""), x.get("cau_hoi") or x.get("van_de", ""), x.get("vi_tri", ""), "", "",
                "Chưa xử lý", ""]
        for j, v in enumerate(vals, 1):
            c = gv.cell(r, j, v)
            c.font = f(11)
            c.border = BOX
            c.alignment = al("center" if j in (1, 2, 4, 6, 7) else "left", "top")
        lv = LEVELS.get(x.get("muc"))
        if lv:
            gv.cell(r, 2).fill = fill(lv["fill"])
            gv.cell(r, 2).font = f(11, True, lv["font"])
        gv.cell(r, 6).number_format = "dd/mm/yyyy"
        dv.add(gv.cell(r, 7))
    fit_rows(gv, r3 + 1, r3 + len(diem), w3)
    gv.freeze_panes = gv.cell(r3 + 1, 1)
    setup_page(gv, paper, title_rows=f"{r3}:{r3}")

    # ---------------- Sheet: Quet tu dong (tuy chon) -------------
    if data.get("ket_qua_quet"):
        qt = wb.create_sheet("Kết quả quét")
        for i in range(1, 7):
            qt.column_dimensions[L(i)].width = 26
        rq = letterhead(qt, 6, date_text, "KẾT QUẢ QUÉT TỰ ĐỘNG BẢNG TÍNH")
        for i, line in enumerate(data["ket_qua_quet"].splitlines()):
            qt.merge_cells(start_row=rq + i, start_column=1, end_row=rq + i, end_column=6)
            c = qt.cell(rq + i, 1, line)
            c.font = f(10, line.startswith("#"))
            c.alignment = al(wrap=True, v="top")
            qt.row_dimensions[rq + i].height = max(15, 14 * (len(line) // 150 + 1))
        setup_page(qt, paper, landscape=True)

    # ---------------- Sheet 1: Tong quan -------------
    tq = wb.create_sheet("Tổng quan", 0)
    NC = 12
    for i in range(1, NC + 1):
        tq.column_dimensions[L(i)].width = 11.5
    r = letterhead(tq, NC, date_text, "KẾT QUẢ RÀ SOÁT HỒ SƠ KHỐI LƯỢNG",
                   data.get("loai_ho_so", "") and f"Hồ sơ {data.get('loai_ho_so','').lower()}")

    # Thong tin chung (2 cot)
    info = [
        ("Dự án", data.get("du_an", "")), ("Gói thầu", data.get("goi_thau", "")),
        ("Hồ sơ rà soát", data.get("ho_so", "")), ("Loại hồ sơ", data.get("loai_ho_so", "")),
        ("Hình thức hợp đồng", data.get("hop_dong", "")), ("Nguồn vốn", data.get("nguon_von", "")),
        ("Tính chất công trình", data.get("tinh_chat", "")), ("Người rà soát", data.get("nguoi_lap", "")),
    ]
    tq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    c = tq.cell(r, 1, "I. THÔNG TIN CHUNG")
    c.font = f(12, True, NAVY)
    r += 1
    for k in range(0, len(info), 2):
        for half, (lab, val) in enumerate(info[k:k + 2]):
            c0 = 1 + half * 6
            tq.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 1)
            tq.merge_cells(start_row=r, start_column=c0 + 2, end_row=r, end_column=c0 + 5)
            a = tq.cell(r, c0, lab)
            a.font = f(11, True)
            a.fill = fill("EEF2F8")
            b = tq.cell(r, c0 + 2, val)
            b.font = f(11)
            for cc in range(c0, c0 + 6):
                tq.cell(r, cc).border = BOX
                tq.cell(r, cc).alignment = al("left")
        tq.row_dimensions[r].height = 20
        r += 1
    r += 1

    # So lieu chinh (the KPI)
    tq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    tq.cell(r, 1, "II. SỐ LIỆU CHÍNH").font = f(12, True, NAVY)
    r += 1
    thieu_count = len(thieu)
    tiles = [
        ("Tổng số điểm cần soát", f"=COUNTA({lvl_rng})", NAVY, "FFFFFF"),
        ("Mức cao", f'=COUNTIF({lvl_rng},"Cao")', "C0392B", "FFFFFF"),
        ("Mức trung bình", f'=COUNTIF({lvl_rng},"Trung bình")', "E6A117", "FFFFFF"),
        ("Mức thấp", f'=COUNTIF({lvl_rng},"Thấp")', "70AD47", "FFFFFF"),
        ("Hạng mục thiếu hồ sơ", thieu_count, "595959", "FFFFFF"),
        ("Hạng mục chủ yếu", data.get("nhom_chu_yeu", ""), "2E75B6", "FFFFFF"),
    ]
    for t, (lab, val, bg, fg) in enumerate(tiles):
        c0 = 1 + t * 2
        tq.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 1)
        tq.merge_cells(start_row=r + 1, start_column=c0, end_row=r + 1, end_column=c0 + 1)
        top = tq.cell(r, c0, val)
        top.font = f(22 if not isinstance(val, str) or val.startswith("=") else 13, True, fg)
        top.alignment = al("center")
        bot = tq.cell(r + 1, c0, lab)
        bot.font = f(10, False, fg)
        bot.alignment = al("center")
        for rr in (r, r + 1):
            for cc in (c0, c0 + 1):
                tq.cell(rr, cc).fill = fill(bg)
                tq.cell(rr, cc).border = Border(left=Side(style="thick", color="FFFFFF"),
                                                right=Side(style="thick", color="FFFFFF"))
    tq.row_dimensions[r].height = 40
    tq.row_dimensions[r + 1].height = 30
    r += 3

    # Danh gia chung
    tq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    tq.cell(r, 1, "III. ĐÁNH GIÁ CHUNG").font = f(12, True, NAVY)
    r += 1
    muc_chung = data.get("danh_gia_chung", "")
    lv = LEVELS.get(muc_chung)
    tq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    c = tq.cell(r, 1, f"Mức {muc_chung.lower()}" if muc_chung else "")
    c.font = f(12, True, lv["font"] if lv else "000000")
    c.fill = fill(lv["fill"] if lv else "FFFFFF")
    c.alignment = al("center")
    tq.merge_cells(start_row=r, start_column=3, end_row=r, end_column=NC)
    c = tq.cell(r, 3, data.get("nhan_xet_chung", ""))
    c.font = f(11)
    c.alignment = al("left")
    for cc in range(1, NC + 1):
        tq.cell(r, cc).border = BOX
    tq.row_dimensions[r].height = max(30, 15 * (len(data.get("nhan_xet_chung", "")) // 110 + 1))
    r += 1
    for lab, key in (("Cần tập trung rà soát", "tap_trung"), ("Cần yêu cầu bổ sung", "can_bo_sung")):
        items = data.get(key) or []
        if not items:
            continue
        tq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        a = tq.cell(r, 1, lab)
        a.font = f(11, True)
        a.fill = fill("EEF2F8")
        a.alignment = al("left", "top")
        tq.merge_cells(start_row=r, start_column=3, end_row=r, end_column=NC)
        b = tq.cell(r, 3, "\n".join(f"- {s}" for s in items))
        b.font = f(11)
        b.alignment = al("left", "top")
        for cc in range(1, NC + 1):
            tq.cell(r, cc).border = BOX
        tq.row_dimensions[r].height = 16 * len(items) + 6
        r += 1
    r += 1

    # Bang so lieu cho bieu do (dat o cot N-P, an khi in)
    tq.column_dimensions["N"].width = 16
    tq.column_dimensions["O"].width = 8
    hr = r
    tq.cell(hr, 14, "Mức rủi ro").font = f(9, True, "808080")
    tq.cell(hr, 15, "Số điểm").font = f(9, True, "808080")
    for i, k in enumerate(LEVELS, 1):
        tq.cell(hr + i, 14, k).font = f(9, color="808080")
        tq.cell(hr + i, 15, f'=COUNTIF({lvl_rng},"{k}")').font = f(9, color="808080")
    parts = list(OrderedDict.fromkeys(x.get("phan", "Khác") or "Khác" for x in diem))
    pr = hr + len(LEVELS) + 2
    tq.cell(pr, 14, "Phần công trình").font = f(9, True, "808080")
    tq.cell(pr, 15, "Số điểm").font = f(9, True, "808080")
    for i, p in enumerate(parts, 1):
        tq.cell(pr + i, 14, p).font = f(9, color="808080")
        tq.cell(pr + i, 15, f'=COUNTIF({part_rng},"{p}")').font = f(9, color="808080")

    tq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    tq.cell(r, 1, "IV. PHÂN BỐ CÁC ĐIỂM CẦN SOÁT").font = f(12, True, NAVY)
    r += 1

    ch1 = BarChart()
    ch1.type = "col"
    ch1.title = "Theo mức rủi ro"
    ch1.style = 10
    ch1.legend = None
    ch1.y_axis.majorGridlines = None
    ch1.add_data(Reference(tq, min_col=15, min_row=hr, max_row=hr + len(LEVELS)), titles_from_data=True)
    ch1.set_categories(Reference(tq, min_col=14, min_row=hr + 1, max_row=hr + len(LEVELS)))
    s = ch1.series[0]
    for i, k in enumerate(LEVELS):
        pt = DataPoint(idx=i)
        pt.graphicalProperties = GraphicalProperties(solidFill=LEVELS[k]["bar"])
        s.dPt.append(pt)
    s.dLbls = no_extra_labels(DataLabelList())
    ch1.gapWidth = 60
    ch1.height, ch1.width = 7, 11.5
    ch1.y_axis.delete = True
    ch1.x_axis.delete = False
    tq.add_chart(ch1, f"A{r}")

    ch2 = BarChart()
    ch2.type = "bar"
    ch2.title = "Theo phần công trình"
    ch2.style = 10
    ch2.legend = None
    ch2.y_axis.majorGridlines = None
    ch2.add_data(Reference(tq, min_col=15, min_row=pr, max_row=pr + max(len(parts), 1)), titles_from_data=True)
    ch2.set_categories(Reference(tq, min_col=14, min_row=pr + 1, max_row=pr + max(len(parts), 1)))
    ch2.series[0].graphicalProperties = GraphicalProperties(solidFill="2E75B6")
    ch2.series[0].dLbls = no_extra_labels(DataLabelList())
    ch2.gapWidth = 50
    ch2.height, ch2.width = 7, 14.5
    ch2.y_axis.delete = True
    ch2.x_axis.delete = False
    tq.add_chart(ch2, f"F{r}")
    r += 15

    # 5 diem uu tien
    tq.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    tq.cell(r, 1, "V. CÁC ĐIỂM CẦN XỬ LÝ TRƯỚC").font = f(12, True, NAVY)
    r += 1
    spans = [(1, 1, "STT"), (2, 2, "Mức"), (3, 5, "Hạng mục"), (6, 12, "Vấn đề cần làm rõ")]
    for a, b, h in spans:
        if b > a:
            tq.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        c = tq.cell(r, a, h)
        c.font = f(11, True, "FFFFFF")
        for cc in range(a, b + 1):
            tq.cell(r, cc).fill = fill(NAVY)
            tq.cell(r, cc).border = BOX
        c.alignment = al("center")
    r += 1
    for i, x in enumerate(diem[:5], 1):
        vals = [i, x.get("muc", ""), x.get("hang_muc", ""), x.get("van_de", "")]
        for (a, b, _), v in zip(spans, vals):
            if b > a:
                tq.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
            c = tq.cell(r, a, v)
            c.font = f(11)
            c.alignment = al("center" if a <= 2 else "left")
            for cc in range(a, b + 1):
                tq.cell(r, cc).border = BOX
        lv = LEVELS.get(x.get("muc"))
        if lv:
            tq.cell(r, 2).fill = fill(lv["fill"])
            tq.cell(r, 2).font = f(11, True, lv["font"])
        tq.row_dimensions[r].height = max(22, 15 * (len(x.get("van_de", "")) // 75 + 1))
        r += 1

    # Chu ky (khi can ky)
    if data.get("chu_ky"):
        r += 1
        sig = data.get("chu_ky_ben") or [["NGƯỜI RÀ SOÁT", data.get("nguoi_lap", "")],
                                          ["TRƯỞNG PHÒNG KỸ THUẬT", "Hoàng Đức Vũ"]]
        span = NC // len(sig)
        for i, (title, name) in enumerate(sig):
            a = 1 + i * span
            b = a + span - 1
            t = (title or "").strip()
            l1, l2 = t, ""
            for pre in ("PHÓ TRƯỞNG PHÒNG", "TRƯỞNG PHÒNG"):      # "TRƯỞNG PHÒNG" trên, tên phòng dưới
                if t.upper().startswith(pre + " "):
                    l1, l2 = t[:len(pre)], t[len(pre):].strip()
                    break
            for rr, txt, ft in ((r, l1, f(11, True)), (r + 1, l2, f(11, True)),
                                (r + 2, "(Ký, ghi rõ họ tên)", f(10, italic=True)),
                                (r + 6, name, f(11, True))):
                tq.merge_cells(start_row=rr, start_column=a, end_row=rr, end_column=b)
                c = tq.cell(rr, a, txt)
                c.font = ft
                c.alignment = al("center", wrap=False)
        r += 7

    setup_page(tq, paper, landscape=False)
    tq.print_area = f"A1:{L(NC)}{r}"
    for col in "NO":
        tq.column_dimensions[col].hidden = False

    wb.active = 0
    wb.save(out)
    print(f"Đã tạo {out}: {len(diem)} điểm, {len(thieu)} hạng mục thiếu hồ sơ, khổ {paper}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("Dùng: python3 qkl_report.py <ket_qua.json> <output.xlsx>")
    with open(sys.argv[1], encoding="utf-8") as fh:
        build(json.load(fh), sys.argv[2])
