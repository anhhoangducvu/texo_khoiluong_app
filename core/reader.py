# -*- coding: utf-8 -*-
"""Đọc hồ sơ khối lượng: Excel (.xlsx/.xlsm/.xls) và PDF có lớp chữ.

Mọi nguồn được đưa về cùng một dạng `Sheet`:
  grid        : list các dòng, mỗi dòng là list giá trị ô (dùng để đọc số liệu)
  head_grid   : như grid nhưng ô gộp (merge) vùng tiêu đề được điền lại (dùng nhận diện tiêu đề)
  formula     : list các dòng, True nếu ô là công thức (chỉ Excel)
  hidden_rows : tập chỉ số dòng (0-based) đang bị ẩn
  labels      : nhãn vị trí từng dòng để ghi vào kết quả ("dòng 15", "tr.3")
"""
import os
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field

ERR_PREFIX = ("#REF", "#VAL", "#DIV", "#N/A", "#NAM", "#NUM", "#NUL")


@dataclass
class Sheet:
    name: str
    kind: str                      # "excel" | "pdf"
    grid: list
    head_grid: list
    formula: list = None
    hidden_rows: set = field(default_factory=set)
    hidden: bool = False
    labels: list = None
    errors: list = field(default_factory=list)     # ["F12", ...]

    def label(self, i):
        if self.labels and i < len(self.labels):
            return self.labels[i]
        return f"dòng {i + 1}"


@dataclass
class Book:
    source: str
    kind: str
    sheets: list
    external_links: int = 0
    warnings: list = field(default_factory=list)
    scanned_pdf: bool = False


# --------------------------------------------------------------------------- Excel
def _recalc_with_soffice(path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    out = tempfile.mkdtemp(prefix="qkl_recalc_")
    try:
        subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", out, path],
                       capture_output=True, timeout=180)
        cand = os.path.join(out, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
        return cand if os.path.exists(cand) else None
    except Exception:
        return None


def _read_xlsx(path, display_name):
    import openpyxl
    wb_f = openpyxl.load_workbook(path, data_only=False)
    wb_v = openpyxl.load_workbook(path, data_only=True)
    warnings = []

    # Ô công thức chưa có giá trị lưu sẵn (file do phần mềm khác xuất) → thử tính lại bằng LibreOffice
    missing = 0
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    if wv[c.coordinate].value is None:
                        missing += 1
                        if missing > 20:
                            break
            if missing > 20:
                break
    if missing > 20:
        rc = _recalc_with_soffice(path)
        if rc:
            wb_v = openpyxl.load_workbook(rc, data_only=True)
            warnings.append("File chưa lưu sẵn kết quả công thức; app đã tính lại bằng LibreOffice.")
        else:
            warnings.append("File có nhiều công thức chưa lưu sẵn kết quả (thường do phần mềm khác xuất). "
                            "Mở file bằng Excel, bấm Lưu rồi tải lại để kết quả quét đầy đủ.")

    ext = 0
    try:
        ext = len(wb_f._external_links or [])
    except Exception:
        pass

    sheets = []
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        nr, nc = ws.max_row or 0, min(ws.max_column or 0, 60)
        grid, formula, errors = [], [], []
        for r in range(1, nr + 1):
            vals, fl = [], []
            for c in range(1, nc + 1):
                v = wv.cell(r, c).value
                fv = ws.cell(r, c).value
                vals.append(v)
                fl.append(isinstance(fv, str) and fv.startswith("="))
                if isinstance(v, str) and v.strip().upper().startswith(ERR_PREFIX):
                    errors.append(ws.cell(r, c).coordinate)
            grid.append(vals)
            formula.append(fl)
        # Điền ô gộp cho vùng tiêu đề (≤ 60 dòng đầu, gộp ≤ 4 dòng)
        head = [list(row) for row in grid[:60]]
        for mr in ws.merged_cells.ranges:
            r1, c1, r2, c2 = mr.min_row, mr.min_col, mr.max_row, mr.max_col
            if r1 > 60 or r2 - r1 > 3:
                continue
            v = grid[r1 - 1][c1 - 1] if r1 - 1 < len(grid) and c1 - 1 < nc else None
            for rr in range(r1, min(r2, 60) + 1):
                for cc in range(c1, min(c2, nc) + 1):
                    if head[rr - 1][cc - 1] is None:
                        head[rr - 1][cc - 1] = v
        head += grid[60:]
        hidden_rows = {r - 1 for r, d in ws.row_dimensions.items() if d.hidden and r <= nr}
        sheets.append(Sheet(name=ws.title, kind="excel", grid=grid, head_grid=head, formula=formula,
                            hidden_rows=hidden_rows, hidden=(ws.sheet_state != "visible"),
                            labels=[f"dòng {i + 1}" for i in range(nr)], errors=errors))
    return Book(source=display_name, kind="excel", sheets=sheets, external_links=ext, warnings=warnings)


def _read_xls(path, display_name):
    import xlrd
    try:
        wb = xlrd.open_workbook(path, formatting_info=True)
        fmt_ok = True
    except Exception:
        wb = xlrd.open_workbook(path)
        fmt_ok = False
    sheets = []
    for sh in wb.sheets():
        grid = []
        for r in range(sh.nrows):
            row = []
            for c in range(min(sh.ncols, 60)):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append("#ERR")
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            grid.append(row)
        head = [list(r) for r in grid]
        hidden = set()
        if fmt_ok:
            for (r1, r2, c1, c2) in sh.merged_cells:
                if r1 > 60 or r2 - r1 > 4:
                    continue
                v = grid[r1][c1] if r1 < len(grid) and c1 < len(grid[r1]) else None
                for rr in range(r1, min(r2, len(head))):
                    for cc in range(c1, min(c2, len(head[rr]))):
                        if head[rr][cc] in (None, ""):
                            head[rr][cc] = v
            for r, info in sh.rowinfo_map.items():
                if getattr(info, "hidden", 0):
                    hidden.add(r)
        sheets.append(Sheet(name=sh.name, kind="excel", grid=grid, head_grid=head, formula=None,
                            hidden_rows=hidden, hidden=bool(getattr(sh, "visibility", 0)),
                            labels=[f"dòng {i + 1}" for i in range(len(grid))]))
    return Book(source=display_name, kind="excel", sheets=sheets,
                warnings=["File .xls (định dạng cũ): không kiểm tra được công thức nhập tay. "
                          "Nên lưu lại thành .xlsx để quét đầy đủ."])


# --------------------------------------------------------------------------- PDF
def _read_pdf(path, display_name):
    import pdfplumber
    blocks = []   # [ncols, rows, labels, pages]
    total_chars = 0
    with pdfplumber.open(path) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            total_chars += len(page.chars)
            tables = page.extract_tables()
            if not tables:
                tables = page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
            for t in tables:
                rows = [[(re.sub(r"\(cid:\d+\)", "", c).strip() if isinstance(c, str) else c) for c in row]
                        for row in t if row]
                rows = [r for r in rows if any(x not in (None, "") for x in r)]
                if not rows:
                    continue
                nc = max(len(r) for r in rows)
                rows = [r + [None] * (nc - len(r)) for r in rows]
                labels = [f"tr.{pno} d.{k}" for k in range(1, len(rows) + 1)]
                if blocks and blocks[-1][0] == nc and pno - blocks[-1][3][-1] <= 1:
                    blocks[-1][1].extend(rows)
                    blocks[-1][2].extend(labels)
                    if blocks[-1][3][-1] != pno:
                        blocks[-1][3].append(pno)
                else:
                    blocks.append([nc, rows, labels, [pno]])
    book = Book(source=display_name, kind="pdf", sheets=[])
    if total_chars < 50:
        book.scanned_pdf = True
        book.warnings.append("PDF không có lớp chữ (bản scan). Bản này chưa đọc được ảnh – "
                             "đề nghị xin file Excel gốc hoặc PDF xuất trực tiếp từ Excel.")
        return book
    for i, (nc, rows, labels, pages) in enumerate(blocks, 1):
        pg = f"tr.{pages[0]}" + (f"–{pages[-1]}" if len(pages) > 1 else "")
        head = [list(r) for r in rows]
        # Ô gộp ngang ở vùng tiêu đề PDF chỉ có chữ ở ô đầu: điền sang phải nếu dòng dưới có chữ
        for r in range(min(len(head) - 1, 10)):
            for c in range(1, nc):
                if head[r][c] in (None, "") and head[r][c - 1] not in (None, "") and head[r + 1][c] not in (None, ""):
                    head[r][c] = head[r][c - 1]
        book.sheets.append(Sheet(name=f"PDF bảng {i} ({pg})", kind="pdf", grid=rows, head_grid=head, labels=labels))
    if not book.sheets:
        book.warnings.append("Không tách được bảng nào trong PDF.")
    return book


# --------------------------------------------------------------------------- API
def load(file_or_path, name=None):
    """Đọc file từ đường dẫn hoặc đối tượng upload (có .read()/.name). Trả Book."""
    if hasattr(file_or_path, "read"):
        name = name or getattr(file_or_path, "name", "ho_so")
        data = file_or_path.read()
        suffix = os.path.splitext(name)[1].lower()
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
        tmp.write(data)
        tmp.close()
        path = tmp.name
    else:
        path = str(file_or_path)
        name = name or os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path, name)
    if ext == ".xls":
        return _read_xls(path, name)
    if ext == ".pdf":
        return _read_pdf(path, name)
    raise ValueError(f"Chưa hỗ trợ định dạng {ext}. Dùng .xlsx, .xlsm, .xls hoặc .pdf.")
