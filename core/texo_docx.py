# -*- coding: utf-8 -*-
"""Bộ dựng Word theo mẫu trình bày TEXO (bản Python của texo_docx_kit.js, skill texo-mau-trinh-bay).

Times New Roman · A4 · header 3 cột (logo | công ty + PHÒNG KỸ THUẬT | tên tài liệu)
· footer "Tài liệu nội bộ · … · Trang X / Y" · bảng màu đã duyệt · thẻ KPI · hộp kết luận · khối ký.
"""
import os

from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips, Cm

ORG = {"company": "CÔNG TY CỔ PHẦN TEXO TƯ VẤN VÀ ĐẦU TƯ", "dept": "PHÒNG KỸ THUẬT",
       "signer": "Hoàng Đức Vũ", "signer_title": "Trưởng phòng Kỹ thuật", "city": "Hà Nội"}
C = {"navy": "1F3864", "blue": "2E5496", "label": "D9E1F2", "first": "EAF0F7", "zebra": "F7F9FC",
     "grid": "BFBFBF", "grey": "595959", "white": "FFFFFF",
     "red": "C00000", "redBg": "FDE4E1", "orange": "C45911", "orangeBg": "FFF1DE",
     "green": "2E7D32", "greenBg": "E4F3E6", "teal": "1F6F78", "tealBg": "E3F1F2"}
# Mức rủi ro của rà soát khối lượng → icon + màu của bộ mẫu
SEV = {"Cao": ("🔴", C["red"], C["redBg"]), "Trung bình": ("🟠", C["orange"], C["orangeBg"]),
       "Thấp": ("🔵", C["blue"], C["label"])}
PAGE = {"A4P": dict(w=11906, h=16838, land=False, content=9026),
        "A4L": dict(w=11906, h=16838, land=True, content=13958)}
FONT, EMOJI = "Times New Roman", "Segoe UI Emoji"
LOGO = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "assets", "logo_texo.png")


def rgb(h):
    return RGBColor.from_string(h)


# ----------------------------------------------------------------- chữ
def run(par, text, size=12, b=False, i=False, c=None, emoji=False):
    r = par.add_run(str(text))
    f = EMOJI if emoji else FONT
    r.font.name, r.font.size, r.font.bold, r.font.italic = f, Pt(size), b, i
    if c:
        r.font.color.rgb = rgb(c)
    rpr = r._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), f)
    return r


def para(container, text="", size=12, b=False, i=False, c=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         before=0, after=6, line=1.15, indent=None, keep_next=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), line
    pf.keep_with_next = keep_next
    if indent is not None:
        pf.left_indent = Twips(indent)
    p.alignment = align
    if text != "":
        run(p, text, size, b, i, c)
    return p


def _pborder(p, edge, color, sz=8, space=2):
    ppr = p._p.get_or_add_pPr()
    bd = ppr.find(qn("w:pBdr"))
    if bd is None:
        bd = OxmlElement("w:pBdr")
        ppr.append(bd)
    e = OxmlElement(f"w:{edge}")
    for k, v in (("w:val", "single"), ("w:sz", str(sz)), ("w:space", str(space)), ("w:color", color)):
        e.set(qn(k), v)
    bd.append(e)


def h1(doc, text):
    p = para(doc, text, 14, b=True, c=C["navy"], align=WD_ALIGN_PARAGRAPH.LEFT, before=14, after=7, keep_next=True)
    _pborder(p, "bottom", C["blue"], 8, 2)
    return p


def h2(doc, text):
    return para(doc, text, 12.5, b=True, c=C["blue"], align=WD_ALIGN_PARAGRAPH.LEFT, before=10, after=5, keep_next=True)


def bullet(doc, text, size=11.5, prefix="•"):
    p = para(doc, "", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, line=1.2)
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Twips(400), Twips(-260)
    run(p, prefix + "  ", size, b=True, c=C["blue"])
    run(p, text, size)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ----------------------------------------------------------------- ô, bảng
def _tc_borders(cell, spec):
    """spec: {edge: (sz, color) | None}"""
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:tcBorders"))
    if old is not None:
        tcpr.remove(old)
    bd = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        v = spec.get(edge)
        if v:
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(v[0]))
            e.set(qn("w:color"), v[1])
        else:
            e.set(qn("w:val"), "nil")
        bd.append(e)
    tcpr.append(bd)


GRID = {e: (4, C["grid"]) for e in ("top", "left", "bottom", "right")}
NONE = {}


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear")
    s.set(qn("w:color"), "auto")
    s.set(qn("w:fill"), fill)
    tcpr.append(s)


def margins(cell, top=60, bottom=60, left=100, right=100):
    tcpr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    tcpr.append(m)


def new_table(container, widths, rows=1):
    try:
        t = container.add_table(rows=rows, cols=len(widths))
    except TypeError:                      # header/footer cần tham số width
        t = container.add_table(rows, len(widths), Twips(sum(widths)))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblpr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(sum(widths)))
    tw.set(qn("w:type"), "dxa")
    old = tblpr.find(qn("w:tblW"))
    if old is not None:
        tblpr.remove(old)
    tblpr.append(tw)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for r in t.rows:
        for c, w in zip(r.cells, widths):
            c.width = Twips(w)
    return t


def no_split(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    trpr.append(e)
    if header:
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        trpr.append(h)


def cp(cell, text="", size=10.5, b=False, i=False, c=None, align=WD_ALIGN_PARAGRAPH.LEFT, after=1.5, first=True):
    p = cell.paragraphs[0] if first and not cell.paragraphs[0].runs and len(cell.paragraphs) == 1 else cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after, p.paragraph_format.space_before, p.paragraph_format.line_spacing = Pt(after), Pt(0), 1.1
    if text != "":
        run(p, text, size, b, i, c)
    return p


def sev_cell(cell, level, size=9.5):
    icon, fg, bg = SEV.get(level, ("", C["grey"], C["first"]))
    p = cp(cell, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, icon + " ", size, emoji=True)
    run(p, level, size, b=True, c=fg)
    shade(cell, bg)


def data_table(container, widths, head, rows, size=10, sev_col=None, center_cols=(), zebra=True):
    t = new_table(container, widths, rows=1 + len(rows))
    hdr = t.rows[0]
    no_split(hdr, header=True)
    for j, h in enumerate(head):
        c = hdr.cells[j]
        _tc_borders(c, GRID)
        shade(c, C["blue"])
        margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cp(c, h, size, b=True, c=C["white"], align=WD_ALIGN_PARAGRAPH.CENTER)
    for i, vals in enumerate(rows, 1):
        r = t.rows[i]
        no_split(r)
        for j, v in enumerate(vals):
            c = r.cells[j]
            _tc_borders(c, GRID)
            margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if j in center_cols else WD_CELL_VERTICAL_ALIGNMENT.TOP
            if j == sev_col:
                sev_cell(c, v, size - 0.5)
                continue
            if j == 0:
                shade(c, C["first"])
            elif zebra and i % 2 == 0:
                shade(c, C["zebra"])
            parts = v if isinstance(v, list) else [v]
            for k, part in enumerate(parts):
                if isinstance(part, dict):
                    cp(c, part.get("text", ""), part.get("size", size), part.get("b", False), part.get("i", False),
                       part.get("c"), WD_ALIGN_PARAGRAPH.CENTER if j in center_cols else WD_ALIGN_PARAGRAPH.LEFT,
                       first=(k == 0))
                else:
                    cp(c, "" if part is None else part, size,
                       align=WD_ALIGN_PARAGRAPH.CENTER if j in center_cols else WD_ALIGN_PARAGRAPH.LEFT, first=(k == 0))
    return t


def meta_table(container, pairs, content, label_ratio=0.3):
    lw = round(content * label_ratio)
    t = new_table(container, [lw, content - lw], rows=len(pairs))
    for (l, v), r in zip(pairs, t.rows):
        no_split(r)
        a, b = r.cells
        for c in (a, b):
            _tc_borders(c, GRID)
            margins(c, 70, 70)
        shade(a, C["label"])
        cp(a, l, 11, b=True)
        cp(b, v, 11)
    return t


def kpi_row(container, items, content, value_size=22):
    """items: [(giá trị, nhãn, dòng phụ, màu)] — màu: red|orange|green|blue|navy|teal"""
    pal = {"red": (C["red"], C["redBg"]), "orange": (C["orange"], C["orangeBg"]), "green": (C["green"], C["greenBg"]),
           "blue": (C["blue"], C["label"]), "navy": (C["navy"], C["first"]), "teal": (C["teal"], C["tealBg"])}
    gap, n = 140, len(items)
    w = (content - gap * (n - 1)) // n
    widths = []
    for k in range(n):
        widths.append(w)
        if k < n - 1:
            widths.append(gap)
    widths[-1] += content - sum(widths)
    t = new_table(container, widths)
    no_split(t.rows[0])
    for k, (val, lbl, sub, color) in enumerate(items):
        fg, bg = pal.get(color, pal["blue"])
        c = t.rows[0].cells[k * 2]
        _tc_borders(c, {"top": (24, fg)})
        shade(c, bg)
        margins(c, 80, 80)
        cp(c, val, value_size, b=True, c=fg, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        cp(c, lbl, 10.5, b=True, c=C["navy"], align=WD_ALIGN_PARAGRAPH.CENTER, after=1, first=False)
        cp(c, sub or "", 9, i=True, c=C["grey"], align=WD_ALIGN_PARAGRAPH.CENTER, after=2, first=False)
        if k < n - 1:
            _tc_borders(t.rows[0].cells[k * 2 + 1], NONE)
    return t


def callout(container, title, body_lines, content, color="blue"):
    pal = {"blue": (C["blue"], C["first"]), "green": (C["green"], C["greenBg"]),
           "orange": (C["orange"], C["orangeBg"]), "red": (C["red"], C["redBg"]), "navy": (C["navy"], C["label"])}
    fg, bg = pal[color]
    t = new_table(container, [content])
    no_split(t.rows[0])
    c = t.rows[0].cells[0]
    _tc_borders(c, {"left": (36, fg)})
    shade(c, bg)
    margins(c, 120, 120, 160, 140)
    cp(c, title, 11.5, b=True, c=fg, after=3)
    for line in body_lines:
        p = cp(c, line, 11.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=2, first=False)
        p.paragraph_format.line_spacing = 1.2
    return t


def spacer(doc, pts=6):
    return para(doc, "", after=pts)


# ----------------------------------------------------------------- trang, header, footer
def _field(par, instr, size=8, c=C["grey"]):
    r = par.add_run()
    r.font.name, r.font.size = FONT, Pt(size)
    r.font.color.rgb = rgb(c)
    a = OxmlElement("w:fldChar"); a.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
    z = OxmlElement("w:fldChar"); z.set(qn("w:fldCharType"), "end")
    r._r.append(a); r._r.append(i); r._r.append(z)


def setup_section(sec, key, header_title, header_sub, footer_left):
    pg = PAGE[key]
    sec.orientation = WD_ORIENT.LANDSCAPE if pg["land"] else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (Twips(pg["h"]), Twips(pg["w"])) if pg["land"] else (Twips(pg["w"]), Twips(pg["h"]))
    sec.left_margin = sec.right_margin = Twips(1440)
    sec.top_margin, sec.bottom_margin = Twips(1300), Twips(1080)
    sec.header_distance = sec.footer_distance = Twips(480)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    content = pg["content"]
    # Header 3 cột
    hdr = sec.header
    for p in list(hdr.paragraphs):
        p._p.getparent().remove(p._p)
    for t in list(hdr.tables):
        t._tbl.getparent().remove(t._tbl)
    logo_w = 1250
    lw = round(content * 0.58) - logo_w
    rw = content - logo_w - lw
    t = new_table(hdr, [logo_w, lw, rw])
    cells = t.rows[0].cells
    bottom = {"bottom": (12, C["blue"])}
    for c in cells:
        _tc_borders(c, bottom)
        margins(c, 0, 40)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    if os.path.exists(LOGO):
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(LOGO, width=Cm(1.77), height=Cm(0.68))
    cp(cells[1], ORG["company"], 8, b=True, c=C["navy"], after=0)
    cp(cells[1], ORG["dept"], 8, b=True, c=C["blue"], after=2, first=False)
    cp(cells[2], header_title, 8.5, b=True, c=C["navy"], align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    cp(cells[2], header_sub or "", 8, i=True, c=C["grey"], align=WD_ALIGN_PARAGRAPH.RIGHT, after=2, first=False)
    hp = hdr.add_paragraph()
    hp.paragraph_format.space_after = Pt(0)
    # Footer: bảng 2 ô (chữ trái | số trang phải) – không phụ thuộc tab của kiểu Footer
    ft = sec.footer
    for p in list(ft.paragraphs):
        p._p.getparent().remove(p._p)
    for t0 in list(ft.tables):
        t0._tbl.getparent().remove(t0._tbl)
    pw = 1500
    ftab = new_table(ft, [content - pw, pw])
    fl, fr = ftab.rows[0].cells
    for c in (fl, fr):
        _tc_borders(c, {"top": (4, C["grid"])})
        margins(c, 50, 0, 0, 0)
    cp(fl, footer_left, 8, i=True, c=C["grey"], after=0)
    pr = cp(fr, "Trang ", 8, c=C["grey"], align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    _field(pr, "PAGE")
    run(pr, " / ", 8, c=C["grey"])
    _field(pr, "NUMPAGES")
    fp = ft.add_paragraph()
    fp.paragraph_format.space_after = Pt(0)
    fp.paragraph_format.line_spacing = 0.5
    return content


def add_section(doc, key, header_title, header_sub, footer_left):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    return setup_section(sec, key, header_title, header_sub, footer_left)


def title_lines(title):
    """Chức danh ký: "TRƯỞNG PHÒNG" ở trên, tên phòng ở dưới (anh Vũ chốt 21/10/2026)."""
    t = (title or "").strip()
    for pre in ("PHÓ TỔNG GIÁM ĐỐC", "TỔNG GIÁM ĐỐC", "PHÓ TRƯỞNG PHÒNG", "TRƯỞNG PHÒNG", "PHÓ GIÁM ĐỐC", "GIÁM ĐỐC"):
        if t.upper().startswith(pre + " "):
            return [t[:len(pre)], t[len(pre):].strip()]
    return [t]


def signature_block(doc, content, date, signers, noi_nhan):
    """Nơi nhận | ngày + các người ký (mặc định Người rà soát, Trưởng phòng Kỹ thuật).
    Mỗi phần (chức danh 2 dòng, ghi chú ký, khoảng ký, họ tên) là một hàng để các cột thẳng hàng."""
    lw = round(content * 0.36)
    n = len(signers)
    sw = (content - lw) // n
    widths = [lw] + [sw] * n
    widths[-1] += content - sum(widths)
    t = new_table(doc, widths, rows=6)
    for r in t.rows:
        no_split(r)
        for c in r.cells:
            _tc_borders(c, NONE)
            margins(c, 0, 0)
    left = t.cell(0, 0).merge(t.cell(5, 0))
    cp(left, "Nơi nhận:", 11, b=True, i=True)
    for x in noi_nhan:
        cp(left, "- " + x, 10, first=False)
    top = t.cell(0, 1)
    for k in range(2, n + 1):
        top = top.merge(t.cell(0, k))
    cp(top, f"{ORG['city']}, ngày {date.day:02d} tháng {date.month:02d} năm {date.year}", 11, i=True,
       align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    for k, (title, name) in enumerate(signers, 1):
        lines = title_lines(title)
        cp(t.cell(1, k), lines[0], 11.5, b=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        cp(t.cell(2, k), lines[1] if len(lines) > 1 else "", 11.5, b=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        cp(t.cell(3, k), "(Ký, ghi rõ họ tên)", 10, i=True, c=C["grey"], align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
        cp(t.cell(4, k), "", 11, after=48)
        cp(t.cell(5, k), name or "", 12, b=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    return t
