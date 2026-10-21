# -*- coding: utf-8 -*-
"""TEXO – Rà soát nhanh hồ sơ khối lượng (Phòng Kỹ thuật).

Lớp 1 (luôn chạy, cục bộ, miễn phí): đọc Excel/PDF → quét theo bộ quy tắc JSON dùng chung với skill.
Lớp 2 (tùy chọn): người dùng tự nhập API key → AI rà soát bổ sung trên nền kết quả quét.
Chạy: streamlit run app.py
"""
import copy
import datetime
import hashlib
import io
import json
import os
import tempfile

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

from core import reader, engine, llm, ai
from core import report_excel, report_word
from core.table import COL_KEYS
from core.text import fmt

APP_DIR = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(APP_DIR, "assets", "logo_texo.png")
SECRETS_FILE = os.path.join(APP_DIR, ".secrets.json")
APP_VERSION = "2.4"
LEVELS = ["Cao", "Trung bình", "Thấp"]
PHAN = ["Đất, hố móng", "Cọc", "Móng, tầng hầm", "Bê tông cốt thép", "Hoàn thiện", "Cơ điện", "Cải tạo",
        "Hạ tầng", "Chung"]
ROLES = ["Tổng hợp", "Chi tiết", "Bỏ qua"]

st.set_page_config(page_title="TEXO · Rà soát khối lượng", page_icon=LOGO if os.path.exists(LOGO) else None,
                   layout="wide")

st.markdown("""
<style>
:root { --texo-indigo:#0C0150; --texo-cyan:#01A0C7; }
/* Chừa chỗ cho thanh công cụ Streamlit để không che tiêu đề */
.block-container { padding-top: 4.2rem !important; }
header[data-testid="stHeader"] { background: rgba(255,255,255,.85); backdrop-filter: blur(4px); }
.texo-bar { border-left: 6px solid var(--texo-cyan); padding: .55rem 1rem; margin: 0 0 1rem 0;
            background: rgba(1,160,199,.07); border-radius: 4px; line-height: 1.5; }
.texo-bar .t { font-size: 1.15rem; font-weight: 700; color: var(--texo-indigo); display: block; }
.texo-bar .s { font-size: .9rem; opacity: .8; display: block; }
.login-box { max-width: 420px; margin: 2rem auto 0 auto; }
div[data-testid="stMetricValue"] { font-size: 1.6rem; }
</style>
""", unsafe_allow_html=True)


# =========================================================================== mật khẩu
def _secret(name, default=""):
    try:
        if name in st.secrets:
            return st.secrets[name]
    except Exception:
        pass
    return os.environ.get(name, default)


DEFAULT_PASSWORD = "texo2026"      # đổi bằng APP_PASSWORD trong Secrets (Cloud) hoặc .streamlit/secrets.toml


def password_gate():
    """Luôn yêu cầu đăng nhập. Mật khẩu lấy từ Secrets/biến môi trường, chưa đặt thì dùng mặc định."""
    if st.session_state.get("auth_ok"):
        return True
    pw = str(_secret("APP_PASSWORD", DEFAULT_PASSWORD) or DEFAULT_PASSWORD)
    _, mid, _ = st.columns([1, 1.2, 1])
    with mid:
        if os.path.exists(LOGO):
            st.image(LOGO, width=170)
        st.markdown("### Rà soát hồ sơ khối lượng")
        st.caption("Phòng Kỹ thuật · Công ty CP TEXO Tư vấn và Đầu tư")
        with st.form("login"):
            v = st.text_input("Mật khẩu", type="password")
            ok = st.form_submit_button("Đăng nhập", type="primary", width="stretch")
        if ok:
            if v == pw:
                st.session_state.auth_ok = True
                st.rerun()
            st.error("Mật khẩu không đúng.")
    return False


if not password_gate():
    st.stop()


# =========================================================================== bộ quy tắc
def rules_state():
    if "rules" not in st.session_state:
        st.session_state.rules = engine.load_rules()
    return st.session_state.rules


def save_rules_to_disk(rules):
    try:
        with open(engine.DEFAULT_RULES, "w", encoding="utf-8") as f:
            json.dump(rules, f, ensure_ascii=False, indent=2)
        return True
    except Exception:
        return False


# =========================================================================== API key
def load_saved_keys():
    try:
        with open(SECRETS_FILE, encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_key(provider, key, base_url):
    d = load_saved_keys()
    d[provider] = {"key": key, "base_url": base_url}
    with open(SECRETS_FILE, "w", encoding="utf-8") as f:
        json.dump(d, f)


def sidebar_ai():
    with st.sidebar:
        if os.path.exists(LOGO):
            st.image(LOGO, width=150)
        st.caption(f"Phòng Kỹ thuật · Rà soát khối lượng · v{APP_VERSION}")
        if st.button("Đăng xuất", width="stretch"):
            st.session_state.clear()
            st.rerun()
        r = rules_state()
        st.caption(f"Bộ quy tắc v{r.get('phien_ban', '?')} · {len(r.get('quy_tac', []))} quy tắc · cập nhật {r.get('cap_nhat', '')}")
        st.divider()
        st.subheader("Bộ não AI (tùy chọn)")
        st.caption("Không nhập key: app vẫn chạy đủ phần quét theo quy tắc, dữ liệu không rời máy.")
        provs = list(llm.PROVIDERS)
        prov = st.selectbox("Nhà cung cấp", provs, format_func=lambda p: llm.PROVIDERS[p]["label"], key="ai_prov")
        saved = load_saved_keys().get(prov, {})
        key = st.text_input("API key", value=saved.get("key", ""), type="password", key=f"ai_key_{prov}")
        base = ""
        if llm.PROVIDERS[prov]["needs_base_url"]:
            base = st.text_input("Base URL (bắt buộc)", value=saved.get("base_url", ""),
                                 placeholder="https://openrouter.ai/api/v1", key=f"ai_base_{prov}")
        st.caption(llm.PROVIDERS[prov]["note"])
        mk = f"models_{prov}"
        if st.button("Lấy danh sách model", width="stretch", disabled=not key):
            ok, res = llm.list_models(prov, key, base or None)
            if ok:
                st.session_state[mk] = res
                if not res:
                    st.warning("Key hợp lệ nhưng không có model nào dùng được.")
            else:
                st.error(res)
        models = st.session_state.get(mk, [])
        model = st.selectbox("Model", models, key=f"ai_model_{prov}") if models else None
        if models and st.button("Kiểm tra kết nối", width="stretch"):
            ok, msg = llm.test_connection(prov, model, key, base or None)
            (st.success if ok else st.error)(("Kết nối được: " if ok else "") + msg)
        with st.expander("Lưu key tại máy này"):
            st.caption("Chỉ lưu trên máy cá nhân. Không dùng chung máy/tài khoản đã lưu key, không đưa app "
                       "kèm key lên mạng công khai. Lộ key do cách dùng là trách nhiệm người dùng, không phải lỗi phần mềm.")
            if st.button("Lưu key", disabled=not key):
                save_key(prov, key, base)
                st.success("Đã lưu vào file .secrets.json cạnh app.")
            if st.button("Xoá key đã lưu"):
                d = load_saved_keys()
                d.pop(prov, None)
                with open(SECRETS_FILE, "w", encoding="utf-8") as f:
                    json.dump(d, f)
                st.info("Đã xoá.")
        if key and model:
            st.success("AI sẵn sàng. Khi bấm rà soát bằng AI, bảng tổng hợp sẽ được gửi tới nhà cung cấp đã chọn.")
        elif key:
            st.warning("Có key nhưng chưa chọn model – bấm 'Lấy danh sách model'.")
        else:
            st.info("Đang chạy bản không dùng AI.")
        return {"provider": prov, "key": key, "base_url": base or None, "model": model,
                "ready": bool(key and model)}


AI = sidebar_ai()


# =========================================================================== tiện ích
def file_hash(up):
    return hashlib.md5(up.getvalue()).hexdigest()


def load_book(up, slot):
    h = file_hash(up)
    if st.session_state.get(f"{slot}_hash") != h:
        with st.spinner(f"Đang đọc {up.name} (file lớn, nhiều sheet có thể mất 10–30 giây)…"):
            st.session_state[slot] = reader.load(io.BytesIO(up.getvalue()), name=up.name)
        st.session_state[f"{slot}_hash"] = h
        st.session_state.pop(f"{slot}_ov", None)
        st.session_state.pop(f"{slot}_an_cache", None)
        if slot == "book":
            for k in ("result", "ai_res", "points_base", "points_edit"):
                st.session_state.pop(k, None)
    return st.session_state[slot]


def col_options(sheet, tmap):
    g = sheet.head_grid
    ncols = max((len(r) for r in g[: (tmap.data_start or 5) + 3]), default=0) if g else 0
    opts = ["—"]
    for c in range(ncols):
        txt = ""
        if tmap.header_row is not None:
            parts = []
            for rr in range(tmap.header_row, (tmap.data_start or tmap.header_row + 1)):
                v = g[rr][c] if rr < len(g) and c < len(g[rr]) else None
                if v not in (None, "") and str(v).strip() not in parts:
                    parts.append(str(v).strip())
            txt = " ".join(parts)
        opts.append(f"{get_column_letter(c + 1)}: {txt[:28]}" if txt else get_column_letter(c + 1))
    return opts


def preview_df(sheet, start, n=10):
    rows = sheet.grid[start: start + n]
    if not rows:
        return pd.DataFrame()
    w = max(len(r) for r in rows)
    data = [[("" if v is None else str(v)) for v in (r + [None] * (w - len(r)))] for r in rows]
    df = pd.DataFrame(data, columns=[get_column_letter(i + 1) for i in range(w)])
    df.index = [sheet.label(start + i) for i in range(len(rows))]
    return df


def analyze_cached(book, ov, slot):
    """Phân tích có lưu đệm theo (file, cấu hình) để mỗi lần bấm không phải quét lại cả chục sheet."""
    key = (st.session_state.get(f"{slot}_hash"), json.dumps(ov, sort_keys=True, ensure_ascii=False))
    cache = st.session_state.setdefault(f"{slot}_an_cache", {})
    if key not in cache:
        cache.clear()
        cache[key] = engine.analyze(book, ov)
    return cache[key]


def mapping_ui(book, slot):
    """Bảng vai trò các sheet (một bảng, không phải mỗi sheet một khối) + chỉnh cột cho từng sheet khi cần."""
    ov = st.session_state.setdefault(f"{slot}_ov", {})
    base = analyze_cached(book, ov, slot)
    auto_role = {sh.name: ("Bỏ qua" if sh.hidden else engine.guess_role(sh, sh._tmap)) for sh in book.sheets}

    def cols_txt(a):
        if a.role == "Bỏ qua" or not a.tmap.cols:
            return ""
        return ", ".join(f"{COL_KEYS[k]}: {get_column_letter(c + 1)}" for k, c in a.tmap.cols.items()
                         if k in COL_KEYS and k not in ("stt", "ma"))

    tbl = pd.DataFrame([{"Sheet": a.name, "Ẩn": "ẩn" if a.hidden else "", "Vai trò": a.role,
                         "Dòng công việc": len(a.items), "Cột app nhận được": cols_txt(a)} for a in base])
    n_sum = sum(1 for a in base if a.role == "Tổng hợp")
    st.caption(f"{len(base)} sheet · {n_sum} bảng tổng hợp được quét theo quy tắc · sheet ẩn mặc định Bỏ qua. "
               "Đổi vai trò ngay trong bảng.")
    ed = st.data_editor(tbl, key=f"{slot}_roles_{st.session_state.get(f'{slot}_hash')}", hide_index=True,
                        width="stretch", height=min(420, 40 + 35 * len(tbl)),
                        disabled=["Sheet", "Ẩn", "Dòng công việc", "Cột app nhận được"],
                        column_config={"Vai trò": st.column_config.SelectboxColumn(options=ROLES, width="small"),
                                       "Ẩn": st.column_config.TextColumn(width="small"),
                                       "Dòng công việc": st.column_config.NumberColumn(width="small")})
    changed = False
    for _, row in ed.iterrows():
        name, role = row["Sheet"], row["Vai trò"]
        cur = (ov.get(name) or {}).get("role")
        if role != auto_role.get(name) and role != cur:
            ov.setdefault(name, {})["role"] = role
            changed = True
        elif role == auto_role.get(name) and cur is not None:
            ov[name].pop("role", None)
            if not ov[name]:
                ov.pop(name)
            changed = True
    if changed:
        st.rerun()

    names = [a.name for a in base if a.role != "Bỏ qua"] + [a.name for a in base if a.role == "Bỏ qua"]
    pick = st.selectbox("Chỉnh dòng tiêu đề, cột cho một sheet (khi app nhận sai)", ["—"] + names,
                        key=f"{slot}_pick")
    if pick != "—":
        sh = next(x for x in book.sheets if x.name == pick)
        a = next(x for x in base if x.name == pick)
        opts = col_options(sh, a.tmap)
        with st.form(f"{slot}_form_{pick}"):
            c2, c3 = st.columns(2)
            hr = c2.number_input("Dòng tiêu đề (số thứ tự)", min_value=1, max_value=max(1, len(sh.grid)),
                                 value=(a.tmap.header_row or 0) + 1)
            ds = c3.number_input("Dòng số liệu đầu tiên", min_value=1, max_value=max(1, len(sh.grid)),
                                 value=(a.tmap.data_start or 1) + 1)
            cols, keys = {}, list(COL_KEYS)
            grid = st.columns(5)
            for i, k in enumerate(keys):
                cur = a.tmap.cols.get(k)
                idx = cur + 1 if cur is not None and cur + 1 < len(opts) else 0
                sel = grid[i % 5].selectbox(COL_KEYS[k], opts, index=idx)
                cols[k] = None if sel == "—" else opts.index(sel) - 1
            st.dataframe(preview_df(sh, max(0, hr - 1), 12), width="stretch", height=280)
            b1, b2 = st.columns(2)
            ok = b1.form_submit_button("Áp dụng cho sheet này", type="primary", width="stretch")
            reset = b2.form_submit_button("Trả về nhận diện tự động", width="stretch")
        if ok:
            o = ov.setdefault(pick, {})
            o.update({"header_row": hr - 1, "data_start": ds - 1, "cols": cols})
            st.rerun()
        if reset:
            role = (ov.get(pick) or {}).get("role")
            ov.pop(pick, None)
            if role:
                ov[pick] = {"role": role}
            st.rerun()
    return ov


def points_frame(chosen, rest):
    rows = []
    for checked, lst in ((True, chosen), (False, rest)):
        for f in lst:
            p = f.to_point() if hasattr(f, "to_point") else dict(f)
            rows.append({"Chọn": checked, "Mức": p["muc"], "Vị trí": p["vi_tri"], "Phần": p["phan"],
                         "Hạng mục": p["hang_muc"], "Khối lượng": p["khoi_luong"], "Tỷ trọng %": p["ty_trong"],
                         "Vấn đề": p["van_de"], "Câu hỏi cho người lập/người kiểm": p["cau_hoi"],
                         "Cách kiểm chứng": p["kiem_chung"], "Hồ sơ còn thiếu": p.get("ho_so_thieu", ""),
                         "Quy tắc": p.get("ma_quy_tac", ""), "Nguồn": p.get("nguon", "Quy tắc")})
    return pd.DataFrame(rows)


def set_base(df):
    st.session_state.points_base = df
    st.session_state.points_ver = st.session_state.get("points_ver", 0) + 1
    st.session_state.pop("points_edit", None)


def _s(v):
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v)


def frame_to_points(df):
    out = []
    for _, r in df.iterrows():
        if not r["Chọn"]:
            continue
        van_de = _s(r["Vấn đề"])
        if r["Nguồn"] == "AI" and not van_de.startswith("[AI]"):
            van_de = "[AI] " + van_de
        tt = r["Tỷ trọng %"]
        out.append({"muc": _s(r["Mức"]) or "Trung bình", "vi_tri": _s(r["Vị trí"]), "phan": _s(r["Phần"]) or "Chung",
                    "hang_muc": _s(r["Hạng mục"]), "khoi_luong": _s(r["Khối lượng"]),
                    "ty_trong": None if tt is None or pd.isna(tt) else float(tt),
                    "van_de": van_de, "cau_hoi": _s(r["Câu hỏi cho người lập/người kiểm"]),
                    "kiem_chung": _s(r["Cách kiểm chứng"]), "ho_so_thieu": _s(r["Hồ sơ còn thiếu"]),
                    "ma_quy_tac": _s(r["Quy tắc"])})
    order = {k: i for i, k in enumerate(LEVELS)}
    out.sort(key=lambda p: (order.get(p["muc"], 9), -(p["ty_trong"] or 0)))
    return out


# =========================================================================== giao diện chính
st.markdown('<div class="texo-bar"><span class="t">RÀ SOÁT NHANH HỒ SƠ KHỐI LƯỢNG</span>'
            '<span class="s">Phòng Kỹ thuật TEXO · bộ lọc trước khi trình ký: chỉ ra 5–20 điểm cần làm rõ, '
            'không tính lại, không soát hình thức.</span></div>', unsafe_allow_html=True)

tab1, tab2, tab3 = st.tabs(["① Rà soát hồ sơ", "② Bộ quy tắc", "③ Hướng dẫn"])

def render_review():
    c1, c2 = st.columns([1.4, 1])
    up = c1.file_uploader("Hồ sơ khối lượng (Excel .xlsx/.xlsm/.xls hoặc PDF có lớp chữ)",
                          type=["xlsx", "xlsm", "xls", "pdf"])
    up_prev = c2.file_uploader("Hồ sơ đợt trước (tùy chọn, để đối chiếu lũy kế)", type=["xlsx", "xlsm", "xls", "pdf"])

    with st.expander("Thông tin hồ sơ và đầu ra", expanded=not bool(up)):
        a1, a2, a3 = st.columns(3)
        info = {
            "du_an": a1.text_input("Dự án"),
            "goi_thau": a2.text_input("Gói thầu"),
            "loai_ho_so": a3.selectbox("Loại hồ sơ", ["Thanh toán", "Nghiệm thu", "Quyết toán", "Hoàn công", "Dự toán"]),
            "hop_dong": a1.selectbox("Hình thức hợp đồng", ["", "Trọn gói", "Đơn giá cố định", "Đơn giá điều chỉnh", "Theo thời gian"]),
            "nguon_von": a2.selectbox("Nguồn vốn", ["", "Vốn nhà nước", "Vốn ngoài ngân sách", "Vốn FDI", "Vốn tư nhân"]),
            "tinh_chat": a3.selectbox("Tính chất", ["Xây mới", "Cải tạo", "Xây mới + cải tạo"]),
            "nguoi_lap": a1.text_input("Người rà soát"),
            "kho_giay": a2.selectbox("Khổ giấy", ["A4", "A3"]),
        }
        n_points = a3.slider("Số điểm tối đa", 5, 30, int(rules_state()["cai_dat"].get("so_diem_toi_da", 20)))
        info["chu_ky"] = st.checkbox("Kèm báo cáo Word (mẫu trình bày TEXO, A4, có khối ký)", value=False)
        if info["loai_ho_so"] == "Quyết toán":
            st.caption("Quyết toán lập theo cách cộng các đợt: các đợt đã ký tạm chấp nhận, chỉ đối chiếu lũy kế; "
                       "tập trung phần mới và phần thiếu hồ sơ (CH-21).")

    if not up:
        st.info("Tải hồ sơ lên để bắt đầu. Có thể thử với file mẫu trong thư mục tests/mau.")
        return

    book = load_book(up, "book")
    info["ho_so"] = book.source
    for w in book.warnings:
        st.warning(w)
    if book.scanned_pdf or not book.sheets:
        return
    prev_book = load_book(up_prev, "prev") if up_prev else None
    if prev_book:
        for w in prev_book.warnings:
            st.caption("Đợt trước: " + w)

    st.markdown("**Bước 1 – Kiểm tra app đọc bảng đúng chưa** (sai thì chọn lại cột, dòng tiêu đề, vai trò sheet)")
    ov = mapping_ui(book, "book")
    pov = None
    if prev_book:
        with st.expander("Nhận diện bảng hồ sơ đợt trước"):
            pov = mapping_ui(prev_book, "prev")

    if st.button("Bước 2 – Rà soát theo bộ quy tắc", type="primary", width="stretch"):
        rules = rules_state()
        with st.spinner("Đang quét…"):
            analyses, raw, grouped = engine.run_all(book, rules, ov, prev_book, pov)
            chosen, rest = engine.select(grouped, rules["cai_dat"], n_points)
        st.session_state.result = {"analyses": analyses, "chosen": chosen, "rest": rest,
                                   "errors": [f for f in raw if f.nguon == "Lỗi"]}
        set_base(points_frame(chosen, rest))
        st.session_state.pop("ai_res", None)

    res = st.session_state.get("result")
    if not res:
        return

    analyses, chosen, rest = res["analyses"], res["chosen"], res["rest"]
    for e in res["errors"]:
        st.error(f"{e.ma}: {e.van_de}")
    df = st.session_state.points_edit if "points_edit" in st.session_state else st.session_state.points_base
    sel = df[df["Chọn"]]
    summ = [a for a in analyses if a.role == "Tổng hợp" and a.items] or [a for a in analyses if a.items]
    p_txt = "–"
    if summ:
        _, k, share = engine.pareto(max(summ, key=lambda a: a.total_value))
        if k:
            p_txt = f"{k} dòng = {share * 100:.0f}%"
    m = st.columns(6)
    m[0].metric("Điểm đã chọn", len(sel))
    m[1].metric("🔴 Cao", int((sel["Mức"] == "Cao").sum()))
    m[2].metric("🟠 Trung bình", int((sel["Mức"] == "Trung bình").sum()))
    m[3].metric("🟢 Thấp", int((sel["Mức"] == "Thấp").sum()))
    m[4].metric("Dòng công việc", sum(len(a.items) for a in summ))
    m[5].metric("Nhóm chủ yếu", p_txt)

    st.markdown("**Bước 3 – Xem, sửa, chọn điểm đưa vào báo cáo** (bỏ tick để loại; tick các dòng cuối để thêm; "
                "sửa trực tiếp mức, câu chữ, hồ sơ còn thiếu)")
    edited = st.data_editor(
        st.session_state.points_base, key=f"editor_{st.session_state.points_ver}", width="stretch", hide_index=True, height=480,
        column_config={
            "Chọn": st.column_config.CheckboxColumn(width="small"),
            "Mức": st.column_config.SelectboxColumn(options=LEVELS, width="small"),
            "Phần": st.column_config.SelectboxColumn(options=PHAN),
            "Tỷ trọng %": st.column_config.NumberColumn(format="%.1f", width="small"),
            "Vấn đề": st.column_config.TextColumn(width="large"),
            "Câu hỏi cho người lập/người kiểm": st.column_config.TextColumn(width="large"),
            "Nguồn": st.column_config.TextColumn(disabled=True, width="small"),
        })
    st.session_state.points_edit = edited

    with st.expander("Kết quả nhận diện từng sheet và các phát hiện khác"):
        st.text(engine.scan_text(analyses, rest, book))

    # ---------------- AI
    st.markdown("**Bước 4 (tùy chọn) – AI rà soát bổ sung**")
    if not AI["ready"]:
        st.caption("Nhập API key và chọn model ở thanh bên trái để dùng. Không có AI, kết quả trên vẫn dùng được.")
    else:
        st.caption("Bảng tổng hợp (tối đa 160 dòng tỷ trọng lớn nhất) và các điểm đã phát hiện sẽ được gửi tới "
                   f"{llm.PROVIDERS[AI['provider']]['label']}.")
        if st.button("AI rà soát bổ sung", width="stretch"):
            pts = frame_to_points(edited)
            with st.spinner("AI đang đọc hồ sơ (20–90 giây)…"):
                r = ai.review(AI["provider"], AI["model"], AI["key"], AI["base_url"], info, pts, analyses,
                              rules_state())
            st.session_state.ai_res = r
            if r.get("ok"):
                add = pd.DataFrame([{**{"Chọn": True}, **{
                    "Mức": p.get("muc"), "Vị trí": p.get("vi_tri", ""), "Phần": p.get("phan") if p.get("phan") in PHAN else "Chung",
                    "Hạng mục": p.get("hang_muc", ""), "Khối lượng": p.get("khoi_luong", ""), "Tỷ trọng %": None,
                    "Vấn đề": p.get("van_de", ""), "Câu hỏi cho người lập/người kiểm": p.get("cau_hoi", ""),
                    "Cách kiểm chứng": p.get("kiem_chung", ""), "Hồ sơ còn thiếu": p.get("ho_so_thieu", ""),
                    "Quy tắc": p.get("ma_quy_tac", ""), "Nguồn": "AI"}} for p in r["data"]["diem_bo_sung"]])
                if len(add):
                    base_df = edited[edited["Nguồn"] != "AI"]
                    set_base(pd.concat([add, base_df], ignore_index=True))
                st.rerun()
        r = st.session_state.get("ai_res")
        if r:
            if not r.get("ok"):
                st.error(r.get("error"))
                if r.get("raw"):
                    st.code(r["raw"][:2000])
            else:
                d = r["data"]
                if r.get("truncated"):
                    st.warning("Phản hồi của AI có thể bị cắt cụt; app đã tự vá JSON.")
                st.success(f"AI bổ sung {len(d['diem_bo_sung'])} điểm (đánh dấu Nguồn = AI ở đầu bảng).")
                if d.get("nhan_xet_chung"):
                    st.markdown("Nhận xét chung của AI:")
                    st.info(d["nhan_xet_chung"])
                if d["danh_gia_lai"]:
                    st.markdown("AI đề nghị đổi mức (anh em tự sửa ở bảng nếu đồng ý):")
                    for x in d["danh_gia_lai"]:
                        st.write(f"- Điểm số {x.get('stt')}: {x.get('muc')} – {x.get('ly_do', '')}")
                if d["goi_y_quy_tac_moi"]:
                    st.markdown("Gợi ý quy tắc mới (xem xét rồi thêm ở tab Bộ quy tắc):")
                    for x in d["goi_y_quy_tac_moi"]:
                        st.write(f"- [{x.get('muc', '')}] {x.get('ten', '')}: {x.get('dau_hieu', '')} → {x.get('cau_hoi', '')}")

    # ---------------- Xuất file
    st.markdown("**Bước 5 – Xuất kết quả**")
    e1, e2 = st.columns(2)
    nx_ai = (st.session_state.get("ai_res") or {}).get("data", {}).get("nhan_xet_chung", "") \
        if (st.session_state.get("ai_res") or {}).get("ok") else ""
    info["nhan_xet_them"] = e1.text_area("Nhận xét chung (thêm vào đầu phần đánh giá)", value=nx_ai, height=110)
    can = e2.text_area("Tài liệu cần yêu cầu bổ sung (mỗi dòng một tài liệu)", height=110)
    info["can_bo_sung"] = [x.strip() for x in can.splitlines() if x.strip()]
    info["ngay"] = datetime.date.today().isoformat()
    pts = frame_to_points(edited)
    data = engine.build_report(analyses, [], info, rules_state(), book, rest=[])
    data["diem"] = pts
    cnt = {k: sum(1 for p in pts if p["muc"] == k) for k in LEVELS}
    data["danh_gia_chung"] = "Cao" if cnt["Cao"] else ("Trung bình" if cnt["Trung bình"] else "Thấp")
    data["nhan_xet_chung"] = ((info["nhan_xet_them"].strip() + " ") if info["nhan_xet_them"].strip() else "") + \
        f"Rà soát {len(pts)} điểm ({cnt['Cao']} Cao, {cnt['Trung bình']} Trung bình, {cnt['Thấp']} Thấp) theo bộ quy tắc " \
        f"phiên bản {rules_state().get('phien_ban')}" + (", có AI rà soát bổ sung" if any(p['van_de'].startswith('[AI]') for p in pts) else "") + \
        ". Kết quả là gợi ý của bộ lọc, người rà soát đối chiếu hồ sơ trước khi kết luận."
    data["ket_qua_quet"] = engine.scan_text(analyses, rest, book)
    stem = (info["du_an"] or os.path.splitext(book.source)[0]).strip().replace("/", "-")[:60]

    if not pts:
        st.warning("Chưa chọn điểm nào.")
    else:
        b1, b2, b3 = st.columns(3)
        with tempfile.TemporaryDirectory() as td:
            xp = os.path.join(td, "out.xlsx")
            report_excel.build(data, xp)
            b1.download_button("⬇️ Excel rà soát (mẫu TEXO)", open(xp, "rb").read(),
                               file_name=f"{stem} - Ra soat khoi luong.xlsx", width="stretch",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            if info["chu_ky"]:
                wp = os.path.join(td, "out.docx")
                report_word.build(data, wp)
                b2.download_button("⬇️ Báo cáo Word (mẫu TEXO A4)", open(wp, "rb").read(),
                                   file_name=f"{stem} - Bao cao ra soat khoi luong.docx", width="stretch",
                                   mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            else:
                b2.caption("Tick 'Kèm báo cáo Word' ở phần Thông tin để có file Word.")
        b3.download_button("⬇️ JSON kết quả (để gửi Claude/skill)", json.dumps(data, ensure_ascii=False, indent=2),
                           file_name=f"{stem} - ket_qua.json", width="stretch", mime="application/json")




with tab1:
    render_review()


# =========================================================================== tab bộ quy tắc
with tab2:
    rules = rules_state()
    st.markdown(f"**{rules.get('ten', '')}** – phiên bản {rules.get('phien_ban')} · cập nhật {rules.get('cap_nhat')}")
    st.caption("Bộ quy tắc này dùng chung với skill texo-khoiluong-radar. Thêm quy tắc ở đây rồi tải file JSON về: "
               "chép đè vào rules/quy_tac_tu_dong.json của app (GitHub) và gửi Claude để cập nhật skill.")
    tbl = pd.DataFrame([{
        "Bật": r.get("bat", True), "Mã": r["ma"], "Loại": r["loai"], "Mức": r["muc"], "Tên": r.get("ten", ""),
        "Từ khóa": ", ".join(r.get("tu_khoa") or (r.get("tu_so") or {}).get("tu_khoa") or []),
        "Trạng thái": r.get("trang_thai", ""), "Nguồn": r.get("nguon", "")} for r in rules["quy_tac"]])
    ed = st.data_editor(tbl, hide_index=True, width="stretch", height=420, key="rules_editor",
                        column_config={"Bật": st.column_config.CheckboxColumn(width="small"),
                                       "Mức": st.column_config.SelectboxColumn(options=LEVELS, width="small")},
                        disabled=["Mã", "Loại", "Tên", "Từ khóa", "Trạng thái", "Nguồn"])
    if st.button("Áp dụng thay đổi bật/tắt, mức"):
        for r, (_, row) in zip(rules["quy_tac"], ed.iterrows()):
            r["bat"] = bool(row["Bật"])
            r["muc"] = row["Mức"]
        rules["cap_nhat"] = datetime.date.today().isoformat()
        st.success("Đã áp dụng cho phiên này." + (" Đã ghi vào file quy tắc." if save_rules_to_disk(rules) else ""))

    st.divider()
    st.markdown("**Thêm quy tắc mới**")
    loai = st.radio("Loại quy tắc", ["tu_khoa", "don_vi", "chi_so"], horizontal=True,
                    format_func={"tu_khoa": "Từ khóa trong tên công việc", "don_vi": "Đơn vị tính bắt buộc",
                                 "chi_so": "Chỉ số giữa hai nhóm công việc"}.get)
    f1, f2, f3 = st.columns([1, 1, 2])
    prefix = f1.text_input("Tiền tố mã (CH, DAO, XT, CKN…)", value="CH").strip().upper()
    ma = f2.text_input("Mã quy tắc", value=engine.next_code(rules, prefix) if prefix else "")
    ten = f3.text_input("Tên quy tắc")
    g1, g2, g3 = st.columns(3)
    muc = g1.selectbox("Mức gợi ý", LEVELS, index=1)
    phan = g2.selectbox("Phần công trình", [""] + PHAN)
    tt = g3.selectbox("Trạng thái", ["Tham khảo", "Đã xác nhận", "Chờ xác nhận"])
    csv = lambda s: [x.strip() for x in (s or "").split(",") if x.strip()]
    new = {"ma": ma.strip(), "loai": loai, "bat": True, "muc": muc, "phan": phan, "ten": ten.strip(), "trang_thai": tt}
    if loai in ("tu_khoa", "don_vi"):
        h1, h2 = st.columns(2)
        new["tu_khoa"] = csv(h1.text_input("Từ khóa (một trong các từ, cách nhau dấu phẩy)", placeholder="vận chuyển, đổ thải"))
        new["va_tu_khoa"] = csv(h2.text_input("Và phải có đủ các từ (tùy chọn)"))
        new["va_tu_khoa_mot"] = csv(h1.text_input("Và có ít nhất một trong các từ (tùy chọn)"))
        new["tru_tu_khoa"] = csv(h2.text_input("Trừ các trường hợp có từ (tùy chọn)"))
        if loai == "don_vi":
            new["don_vi_hop_le"] = csv(st.text_input("Đơn vị hợp lệ (viết chuẩn: m3, m2, 100m2, tan, kg, m, bo, cai)",
                                                     value="m3"))
            new["van_de"] = "Đơn vị là \"{dvt}\", không đúng đơn vị thường dùng cho công việc này."
        else:
            new["chi_khi_co_kl"] = st.checkbox("Chỉ báo khi dòng có khối lượng đề nghị", value=True)
    else:
        h1, h2 = st.columns(2)
        new["tu_so"] = {"tu_khoa": csv(h1.text_input("Tử số – từ khóa", placeholder="thép")),
                        "va_tu_khoa": csv(h1.text_input("Tử số – phải có đủ", placeholder="sàn")),
                        "tru_tu_khoa": csv(h1.text_input("Tử số – trừ")),
                        "don_vi": csv(h1.text_input("Tử số – đơn vị", value="tan, kg"))}
        new["mau_so"] = {"tu_khoa": csv(h2.text_input("Mẫu số – từ khóa", placeholder="bê tông")),
                         "va_tu_khoa": csv(h2.text_input("Mẫu số – phải có đủ", placeholder="sàn")),
                         "tru_tu_khoa": csv(h2.text_input("Mẫu số – trừ")),
                         "don_vi": csv(h2.text_input("Mẫu số – đơn vị", value="m3"))}
        k1, k2, k3, k4 = st.columns(4)
        new["quy_doi"] = k1.selectbox("Quy đổi tử số về", ["kg", "m2", "m3"])
        mn = k2.text_input("Chỉ số nhỏ nhất (để trống nếu không có)")
        mx = k3.text_input("Chỉ số lớn nhất (để trống nếu không có)")
        new["don_vi_chi_so"] = k4.text_input("Đơn vị chỉ số", value="kg/m3")
        new["min"] = float(mn.replace(",", ".")) if mn.strip() else None
        new["max"] = float(mx.replace(",", ".")) if mx.strip() else None
        new["van_de"] = f"{ten or 'Chỉ số'} {{chi_so}} {new['don_vi_chi_so']}, ngoài khoảng tham khảo {{min}}–{{max}}."
    if loai == "tu_khoa":
        new["van_de"] = st.text_input("Mô tả vấn đề (có thể dùng {tu_khoa})", value="Tên công việc có chữ \"{tu_khoa}\".")
    new["cau_hoi"] = st.text_input("Câu hỏi cho người lập/người kiểm")
    new["kiem_chung"] = st.text_input("Cách kiểm chứng (hồ sơ cần đối chiếu)")
    new["nguon"] = st.text_input("Nguồn (người gửi, dự án, ngày)",
                                 value=f"Phòng Kỹ thuật, {datetime.date.today():%d/%m/%Y}")
    new = {k: v for k, v in new.items() if v not in ([], "") or k in ("phan",)}

    # Thử trên hồ sơ đang mở
    res = st.session_state.get("result")
    if res and (new.get("tu_khoa") or new.get("tu_so")):
        trial = {"cai_dat": rules["cai_dat"], "quy_tac": [dict(new, muc=new.get("muc", "Trung bình"))]}
        errs = engine.validate_rules(trial)
        if not errs:
            r_ = engine.Runner(trial, res["analyses"], st.session_state.get("book"))
            found = r_.run()
            st.caption(f"Thử trên hồ sơ đang mở: {len(found)} dòng khớp" +
                       (": " + "; ".join(f"{f.vi_tri} – {f.hang_muc[:40]}" for f in found[:5]) if found else "."))
    if st.button("Thêm quy tắc", type="primary"):
        errs = engine.validate_rules({"quy_tac": [new]})
        if any(r["ma"] == new["ma"] and r.get("ten") == new.get("ten") for r in rules["quy_tac"]):
            errs.append("Đã có quy tắc cùng mã và tên.")
        if errs:
            for e in errs:
                st.error(e)
        else:
            rules["quy_tac"].append(new)
            engine.register_code(rules, new["ma"])
            rules["cap_nhat"] = datetime.date.today().isoformat()
            ok = save_rules_to_disk(rules)
            st.success(f"Đã thêm {new['ma']}." + (" Đã ghi vào file quy tắc." if ok else
                                                   " Môi trường này không ghi được file – nhớ tải JSON về."))

    st.divider()
    d1, d2 = st.columns(2)
    d1.download_button("⬇️ Tải bộ quy tắc (JSON)", json.dumps(rules, ensure_ascii=False, indent=2),
                       file_name="quy_tac_tu_dong.json", mime="application/json", width="stretch")
    upr = d2.file_uploader("Nạp bộ quy tắc khác (JSON)", type=["json"], key="rules_up")
    if upr is not None and st.session_state.get("rules_up_hash") != file_hash(upr):
        try:
            newr = json.loads(upr.getvalue().decode("utf-8"))
            errs = engine.validate_rules(newr)
            if errs:
                for e in errs[:10]:
                    st.error(e)
            else:
                st.session_state.rules = newr
                st.session_state.rules_up_hash = file_hash(upr)
                save_rules_to_disk(newr)
                st.success(f"Đã nạp bộ quy tắc v{newr.get('phien_ban')} ({len(newr['quy_tac'])} quy tắc).")
                st.rerun()
        except Exception as e:
            st.error(f"File không hợp lệ: {e}")


# =========================================================================== hướng dẫn
with tab3:
    readme = os.path.join(APP_DIR, "README.md")
    if os.path.exists(readme):
        with open(readme, encoding="utf-8") as f:
            st.markdown(f.read())
